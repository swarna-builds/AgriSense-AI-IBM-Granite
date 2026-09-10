import os
import glob
from typing import Any, List, Tuple
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

PPTX_PATH = "your projectpresntation.pptx"

if not os.path.exists(PPTX_PATH):
    raise FileNotFoundError(f"'{PPTX_PATH}' not found in project root folder.")

prs = Presentation(PPTX_PATH)

FONT_NAME = "Arial"
COLOR_PRIMARY = RGBColor(15, 59, 46)     # Deep Forest Green
COLOR_CARD = RGBColor(245, 248, 246)       # Off-white / light mint
COLOR_BORDER = RGBColor(200, 215, 205)    # Accent border
COLOR_TEXT = RGBColor(30, 41, 59)         # Dark slate

def find_image(base_name: str) -> Any:
    """Finds image matching base name across png, jpg, or double extensions."""
    matches = glob.glob(f"{base_name}*")
    return matches[0] if matches else None

def get_text_frame(shape: Any) -> Any:
    """Safely retrieves text_frame without static type checker warnings."""
    return getattr(shape, "text_frame", None)

def reset_slide_content(slide: Any) -> None:
    """Clears slide content shapes while preserving the slide title placeholder."""
    if not slide.shapes:
        return
    title_shape = slide.shapes[0]
    shapes_to_remove = [s for s in slide.shapes if s != title_shape]
    for s in shapes_to_remove:
        sp = getattr(s, "_element", None)
        if sp is not None and sp.getparent() is not None:
            sp.getparent().remove(sp)

def set_slide_title(slide: Any, title_text: str) -> None:
    """Enforces Arial 28pt bold heading styling aligned with the IBM template."""
    title_box = slide.shapes[0]
    tf = get_text_frame(title_box)
    if tf is not None:
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_NAME
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.alignment = PP_ALIGN.LEFT

def create_card(slide: Any, left: Any, top: Any, width: Any, height: Any, title: str, items: List[str]) -> Any:
    """Builds a styled container card adhering to template typography."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(1.5)
    
    tf = get_text_frame(card)
    if tf is not None:
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.25)
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_NAME
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        
        for item in items:
            pi = tf.add_paragraph()
            pi.text = item
            pi.font.name = FONT_NAME
            pi.font.size = Pt(15)
            pi.font.color.rgb = COLOR_TEXT
            pi.space_before = Pt(6)
    return card

# ==============================================================================
# SLIDE 1: IBM University Engagement Project Submission (Index 0)
# ==============================================================================
slide_1 = prs.slides[0]
reset_slide_content(slide_1)
set_slide_title(slide_1, "IBM University Engagement Project Submission")

card_student = slide_1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(7.5), Inches(5.0))
card_student.fill.solid()
card_student.fill.fore_color.rgb = COLOR_CARD
card_student.line.color.rgb = COLOR_BORDER
card_student.line.width = Pt(1.5)

tf_s1 = get_text_frame(card_student)
if tf_s1 is not None:
    tf_s1.word_wrap = True
    tf_s1.margin_left = Inches(0.35)
    tf_s1.margin_top = Inches(0.3)
    
    p0 = tf_s1.paragraphs[0]
    p0.text = "AgriSense AI"
    p0.font.name = FONT_NAME
    p0.font.size = Pt(26)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY
    
    p_sub = tf_s1.add_paragraph()
    p_sub.text = "Autonomous Agronomic Decision & Regional Agromet Advisory Platform"
    p_sub.font.name = FONT_NAME
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_PRIMARY
    p_sub.space_after = Pt(14)
    
    fields: List[Tuple[str, str]] = [
        ("Domain of Project:", "Agriculture"),
        ("Student Name:", "M SWARNA LAKSHMI"),
        ("Email ID:", "swarnalakshmim2006@gmail.com"),
        ("Mobile Number:", "6360306436"),
        ("Institution:", "S J C Institute of Technology"),
        ("Technology Stack:", "IBM Granite (ibm/granite-4-h-small) | Langflow | watsonx.ai")
    ]
    for lbl, val in fields:
        pi = tf_s1.add_paragraph()
        pi.text = f"{lbl} {val}"
        pi.font.name = FONT_NAME
        pi.font.size = Pt(16)
        pi.font.color.rgb = COLOR_TEXT
        pi.space_before = Pt(6)

photo_img = find_image("passport_photo")
if photo_img:
    slide_1.shapes.add_picture(photo_img, Inches(8.7), Inches(1.8), width=Inches(3.8))
else:
    box = slide_1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.7), Inches(1.8), Inches(3.8), Inches(5.0))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_CARD
    box.line.color.rgb = COLOR_BORDER
    tf_b = get_text_frame(box)
    if tf_b is not None:
        tf_b.text = "[passport_photo not found]"

# ==============================================================================
# SLIDE 2: Problem Statement (Index 1)
# ==============================================================================
slide_2 = prs.slides[1]
reset_slide_content(slide_2)
set_slide_title(slide_2, "Problem Statement")

s2_data = [
    ("Unpredictable Epiphytotics", [
        "• Pathogens like Downy Mildew spread rapidly when humidity exceeds 78% RH.",
        "• Visual foliar symptoms appear days after active spore germination.",
        "• Smallholders lack real-time agromet early warning tools."
    ]),
    ("Agrochemical Misapplication", [
        "• Lack of scientific guidance leads to heavy dealer-driven chemical overdosing.",
        "• Causes severe phytotoxicity and chemical-resistant pathogen strains.",
        "• Excessive toxic runoff contaminates soil and rural groundwater."
    ]),
    ("LLM Hallucination Hazards", [
        "• Generic commercial AI tools invent arbitrary spray dosages.",
        "• Lack grounding in verified regional ICAR and KVK standard packages.",
        "• Vernacular barriers lock non-literate farmers out of critical advice."
    ])
]
for i, (title, items) in enumerate(s2_data):
    x_val = float(0.8 + i * 4.0)
    create_card(slide_2, Inches(x_val), Inches(1.8), Inches(3.7), Inches(5.0), title, items)

# ==============================================================================
# SLIDE 3: Proposed Solution (Index 2)
# ==============================================================================
slide_3 = prs.slides[2]
reset_slide_content(slide_3)
set_slide_title(slide_3, "Proposed Solution")

s3_data = [
    ("Autonomous Agromet Sentinel", [
        "• Continuous background polling of live weather telemetry.",
        "• Proactive red alert dispatch before active spore germination.",
        "• Eliminates reliance on manual user query initiation."
    ]),
    ("Langflow Agentic Orchestration", [
        "• Modular visual workflow connecting sensory inputs to LLM agents.",
        "• Seamlessly integrates custom Python vision and weather tools.",
        "• Coordinates ChromaDB RAG retrieval with IBM Granite."
    ]),
    ("ICAR-Grounded Granite AI", [
        "• Powered by IBM Granite (ibm/granite-4-h-small) on watsonx.ai.",
        "• Restricts treatments to verified ICAR/KVK active ingredients.",
        "• Vernacular audio advisory delivery in KA, TE, HI, and EN."
    ])
]
for i, (title, items) in enumerate(s3_data):
    x_val = float(0.8 + i * 4.0)
    create_card(slide_3, Inches(x_val), Inches(1.8), Inches(3.7), Inches(5.0), title, items)

# ==============================================================================
# SLIDE 4: Project Objectives (Index 3)
# ==============================================================================
slide_4 = prs.slides[3]
reset_slide_content(slide_4)
set_slide_title(slide_4, "Project Objectives")

s4_steps = [
    ("Objective 1: Zero-Latency Agromet Monitoring", "Continuously track hyper-local temperature, relative humidity (>78% RH), and precipitation thresholds to predict pathogen outbreaks."),
    ("Objective 2: Multi-Modal Pathology Diagnostics", "Deterministic OpenCV color segmentation to isolate foliar chlorosis, fruit rot craters, and fungal necrotic lesions."),
    ("Objective 3: Langflow Agentic Flow Orchestration", "Build an extensible, modular visual pipeline connecting user input, sensor telemetry, RAG, and IBM Granite."),
    ("Objective 4: Zero-Hallucination Chemical Grounding", "Integrate ICAR/KVK vector retrieval to guarantee calibrated active ingredient recommendations (g/L or ml/L)."),
    ("Objective 5: Multi-Lingual Vernacular Inclusivity", "Deploy multi-tier regional dialect voice generation (Kannada, Telugu, Hindi, English) for smallholder usability.")
]
for i, (stitle, sdesc) in enumerate(s4_steps):
    y_val = float(1.8 + i * 1.02)
    card = slide_4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_val), Inches(11.7), Inches(0.92))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(1.5)
    tf = get_text_frame(card)
    if tf is not None:
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.08)
        p1 = tf.paragraphs[0]
        p1.text = stitle
        p1.font.name = FONT_NAME
        p1.font.size = Pt(19)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.name = FONT_NAME
        p2.font.size = Pt(14)
        p2.font.color.rgb = COLOR_TEXT

# ==============================================================================
# SLIDE 5: Technology Used (Index 4)
# ==============================================================================
slide_5 = prs.slides[4]
reset_slide_content(slide_5)
set_slide_title(slide_5, "Technology Used")

s5_data = [
    ("Langflow Platform & Components", [
        "• Langflow Platform (Agentic visual pipeline)",
        "• Langflow Component Names:",
        "  - Agent Component (Decision reasoning)",
        "  - Chat Input & Chat Output Components",
        "  - Custom Python Tool Component",
        "  - ChromaDB Vector Store Component"
    ]),
    ("IBM Granite & Cloud Platform", [
        "• IBM watsonx.ai Cloud Platform",
        "• IBM Granite Model: ibm/granite-4-h-small",
        "• IBM Cloud IAM Security Authentication",
        "• Strict Agronomic JSON Schema Enforcement",
        "• Hosting Region: Dallas (us-south)"
    ]),
    ("Vision, Storage & Audio Stack", [
        "• ChromaDB Vector Store (ICAR/KVK protocols)",
        "• OpenCV Deterministic Foliar Segmentation",
        "• Streamlit Native Dashboard Interface",
        "• gTTS Regional Edge Text-to-Speech Engine",
        "• Python 3.10+ Asynchronous Architecture"
    ])
]
for i, (title, items) in enumerate(s5_data):
    x_val = float(0.8 + i * 4.0)
    create_card(slide_5, Inches(x_val), Inches(1.8), Inches(3.7), Inches(5.0), title, items)

# ==============================================================================
# SLIDE 6: Langflow component Used (Index 5)
# ==============================================================================
slide_6 = prs.slides[5]
reset_slide_content(slide_6)
set_slide_title(slide_6, "Langflow component Used")

langflow_components = [
    ("1. Chat Input Component", "Ingests farmer queries, targeted crop selection, field location coordinates, and high-resolution foliar disease photography."),
    ("2. Custom Python Component (Agromet Sentinel)", "Non-blocking daemon executing live agrometeorological evaluation against biological danger zones (RH > 78% and rain probability)."),
    ("3. Custom Vision Component (OpenCV Pathology)", "Performs deterministic RGB color segmentation to compute necrotic lesion surface percentages and foliar chlorosis patterns."),
    ("4. Vector Store Component (ChromaDB RAG)", "Retrieves verified ICAR and KVK agronomic protocols to inject calibrated active ingredients, dilution ratios, and safety intervals."),
    ("5. Agent & Chat Output Component (IBM Granite)", "Orchestrates ibm/granite-4-h-small to synthesize inputs into a 4-tab structured advisory and triggers edge vernacular audio playback.")
]
for i, (comp_title, comp_desc) in enumerate(langflow_components):
    y_val = float(1.8 + i * 1.02)
    card = slide_6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(y_val), Inches(11.7), Inches(0.92))
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD
    card.line.color.rgb = COLOR_BORDER
    card.line.width = Pt(1.5)
    tf = get_text_frame(card)
    if tf is not None:
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.08)
        p1 = tf.paragraphs[0]
        p1.text = comp_title
        p1.font.name = FONT_NAME
        p1.font.size = Pt(19)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY
        p2 = tf.add_paragraph()
        p2.text = comp_desc
        p2.font.name = FONT_NAME
        p2.font.size = Pt(14)
        p2.font.color.rgb = COLOR_TEXT

# ==============================================================================
# SLIDE 7: Technical flow diagram- Architecture blueprint (Index 6)
# ==============================================================================
slide_7 = prs.slides[6]
reset_slide_content(slide_7)
set_slide_title(slide_7, "Technical flow diagram- Architecture blueprint")

tiers = [
    ("Tier 1: Presentation & Vernacular UI", [
        "• Responsive Streamlit Dashboard Interface",
        "• Langflow Chat Interface & Flow Runner",
        "• Vernacular Localization (KA / TE / HI / EN)",
        "• Edge Audio Playback (gTTS Synthesis)",
        "• High-Resolution Foliar Image Uploader"
    ]),
    ("Tier 2: Langflow Orchestration & RAG", [
        "• Modular Langflow Agentic Flow Pipeline",
        "• SentinelDaemon: Async Agromet Poll (RH > 78%)",
        "• VisionService: Deterministic OpenCV Masking",
        "• ChromaDB Vector Store: ICAR/KVK Protocols",
        "• Multi-Variable Agronomic Context Fusion"
    ]),
    ("Tier 3: Foundation Model & Cloud", [
        "• IBM watsonx.ai Platform (Dallas Region)",
        "• Foundation Model: ibm/granite-4-h-small",
        "• Strict Agronomic JSON Schema Enforcement",
        "• IBM Cloud IAM Security Token Authentication",
        "• Autonomous Self-Healing Fallback Daemon"
    ])
]
for i, (tier_title, tier_items) in enumerate(tiers):
    x_val = float(0.8 + i * 4.0)
    create_card(slide_7, Inches(x_val), Inches(1.8), Inches(3.7), Inches(5.0), tier_title, tier_items)

# ==============================================================================
# SLIDE 8: Core Implementation & Functional Modules (Index 7)
# ==============================================================================
slide_8 = prs.slides[7]
reset_slide_content(slide_8)
set_slide_title(slide_8, "Core Implementation & Functional Modules")

s8_data = [
    ("SentinelDaemon Module", [
        "• Asynchronous agromet polling daemon.",
        "• Evaluates RH > 78% and rainfall probability.",
        "• Computes composite biological threat levels.",
        "• Dispatches proactive hazard alerts."
    ]),
    ("AgriRAG Vector Service", [
        "• Segmented ICAR agronomic publications.",
        "• Dense semantic vector similarity search.",
        "• Restricts chemicals to vetted active ingredients.",
        "• Injects exact dilution ratios (g/L or ml/L)."
    ]),
    ("Granite Inference Engine", [
        "• watsonx.ai REST client integration.",
        "• Prompts ibm/granite-4-h-small with JSON schema.",
        "• Generates 4-tab structured clinical guidance.",
        "• Handles fallback parsing gracefully."
    ])
]
for i, (title, items) in enumerate(s8_data):
    x_val = float(0.8 + i * 4.0)
    create_card(slide_8, Inches(x_val), Inches(1.8), Inches(3.7), Inches(5.0), title, items)

# ==============================================================================
# SLIDE 9: Project Demonstration & User Interface (Index 8)
# ==============================================================================
slide_9 = prs.slides[8]
reset_slide_content(slide_9)
set_slide_title(slide_9, "Project Demonstration & User Interface")

s9_cards = [
    ("Tab 1: Agromet Risk Sentinel", [
        "• Real-time temperature, humidity, and rain telemetry gauges.",
        "• Color-coded biological hazard alert banners (Red/Yellow/Green).",
        "• Instant vernacular text-to-speech advisory audio playback."
    ]),
    ("Tab 2: Foliar Pathology Vision", [
        "• Dual-view image comparison: Original field photo vs OpenCV mask.",
        "• Necrotic surface area percentage calculation and severity metrics.",
        "• Symptom classification linked directly to pathogen profile."
    ]),
    ("Tab 3: 7-Day Treatment Roadmap", [
        "• Step-by-step chemical and organic spray intervention timeline.",
        "• Exact ICAR chemical dilutions (e.g., Metalaxyl + Mancozeb @ 2g/L).",
        "• Tank mixture safety, wait periods, and protective equipment guidelines."
    ]),
    ("Tab 4: Economic Decision Simulation", [
        "• Financial comparison: Preventative treatment cost vs crop loss.",
        "• Estimated input chemical costs vs net crop yield revenue saved.",
        "• Actionable ROI projections tailored to farm acreage."
    ])
]
for i, (ctitle, citems) in enumerate(s9_cards):
    row = i // 2
    col = i % 2
    x_val = float(0.8 + col * 6.0)
    y_val = float(1.8 + row * 2.6)
    create_card(slide_9, Inches(x_val), Inches(y_val), Inches(5.7), Inches(2.35), ctitle, citems)

# ==============================================================================
# SLIDE 10: IBM watsonx.ai Model Usage (Index 9)
# ==============================================================================
slide_10 = prs.slides[9]
reset_slide_content(slide_10)
set_slide_title(slide_10, "IBM watsonx.ai Model Usage")

usage_img = find_image("watsonx_usage")
if usage_img:
    slide_10.shapes.add_picture(usage_img, Inches(0.8), Inches(1.8), width=Inches(7.2))
else:
    box = slide_10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(7.2), Inches(4.8))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_CARD
    tf_u = get_text_frame(box)
    if tf_u is not None:
        tf_u.text = "[watsonx_usage image not found]"

card_metrics = slide_10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.8), Inches(4.2), Inches(4.8))
card_metrics.fill.solid()
card_metrics.fill.fore_color.rgb = COLOR_CARD
card_metrics.line.color.rgb = COLOR_BORDER
card_metrics.line.width = Pt(1.5)

tf_m = get_text_frame(card_metrics)
if tf_m is not None:
    tf_m.word_wrap = True
    tf_m.margin_left = Inches(0.25)
    tf_m.margin_top = Inches(0.25)
    
    p = tf_m.paragraphs[0]
    p.text = "Inference Telemetry"
    p.font.name = FONT_NAME
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    
    metrics = [
        ("Foundation Model:", "ibm/granite-4-h-small"),
        ("Serving Platform:", "IBM watsonx.ai"),
        ("Hosting Region:", "Dallas (us-south)"),
        ("Tokens Processed:", "591 Tokens"),
        ("Execution Time:", "1.2s Average Latency"),
        ("Output Schema:", "Strict Agronomic JSON")
    ]
    for label, val in metrics:
        p_item = tf_m.add_paragraph()
        p_item.text = f"{label} {val}"
        p_item.font.name = FONT_NAME
        p_item.font.size = Pt(16)
        p_item.font.color.rgb = COLOR_TEXT
        p_item.space_before = Pt(10)

# ==============================================================================
# SLIDE 11: Socio-Economic Impact (Index 10)
# ==============================================================================
slide_11 = prs.slides[10]
reset_slide_content(slide_11)
set_slide_title(slide_11, "Socio-Economic Impact")

s11_data = [
    ("Cost Optimization", [
        "• Eliminates ~35% of redundant routine chemical sprays.",
        "• Times agrochemicals strictly to biological infection windows.",
        "• Maximizes return on smallholder agricultural input costs."
    ]),
    ("Yield Preservation", [
        "• 48-72 hour lead time before active spore propagation.",
        "• Rapid disease suppression preserves marketable crop.",
        "• Protects harvest quality and seasonal farming revenue."
    ]),
    ("Environmental Safety", [
        "• Curtails toxic groundwater chemical leaching.",
        "• Preserves beneficial soil microbiome and pollinators.",
        "• Prevents regional chemical resistance development."
    ])
]
for i, (title, items) in enumerate(s11_data):
    x_val = float(0.8 + i * 4.0)
    create_card(slide_11, Inches(x_val), Inches(1.8), Inches(3.7), Inches(5.0), title, items)

# ==============================================================================
# SLIDE 12: Git Hub Link (Index 11)
# ==============================================================================
slide_12 = prs.slides[11]
reset_slide_content(slide_12)
set_slide_title(slide_12, "Git Hub Link")

card_git = slide_12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
card_git.fill.solid()
card_git.fill.fore_color.rgb = COLOR_CARD
card_git.line.color.rgb = COLOR_BORDER
card_git.line.width = Pt(1.5)

tf_g = get_text_frame(card_git)
if tf_g is not None:
    tf_g.word_wrap = True
    tf_g.margin_left = Inches(0.3)
    tf_g.margin_top = Inches(0.25)
    
    p = tf_g.paragraphs[0]
    p.text = "Repository URL:"
    p.font.name = FONT_NAME
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    
    p_link = tf_g.add_paragraph()
    p_link.text = "https://github.com/swarna-builds/AgriSense-AI-IBM-Granite"
    p_link.font.name = FONT_NAME
    p_link.font.size = Pt(20)
    p_link.font.bold = True
    p_link.font.color.rgb = RGBColor(10, 102, 194)
    p_link.space_after = Pt(14)
    
    p_h2 = tf_g.add_paragraph()
    p_h2.text = "Mandatory Deliverables Verified on Remote Main Branch:"
    p_h2.font.name = FONT_NAME
    p_h2.font.size = Pt(20)
    p_h2.font.bold = True
    p_h2.font.color.rgb = COLOR_PRIMARY
    p_h2.space_before = Pt(8)
    
    delivs = [
        "• app.py & app.json — Langflow agentic workflow manifest and service dependencies",
        "• yourproblemstatement.pdf — Formal problem formulation & ICAR objectives",
        "• your projectpresntation.pptx — Complete technical deck and architectural artifacts",
        "• Source Code Repository — core/, services/, and automated test suite"
    ]
    for d in delivs:
        pd = tf_g.add_paragraph()
        pd.text = d
        pd.font.name = FONT_NAME
        pd.font.size = Pt(18)
        pd.font.color.rgb = COLOR_TEXT
        pd.space_before = Pt(6)

# ==============================================================================
# SLIDE 13: Future Scope (Index 12)
# ==============================================================================
slide_13 = prs.slides[12]
reset_slide_content(slide_13)
set_slide_title(slide_13, "Future Scope")

s13_data = [
    ("Phase 1: Delivered Platform", [
        "• Deployed Granite-4-H-Small agronomic engine.",
        "• Validated ICAR-grounded vector RAG retrieval.",
        "• Multi-modal Langflow agentic workflow.",
        "• Quad-language vernacular text-to-speech."
    ]),
    ("Phase 2: IoT Sensor Ingestion", [
        "• Integrate LoRaWAN field probes for soil N-P-K.",
        "• Ingest live leaf wetness and soil moisture levels.",
        "• Micro-climate prediction models at farm level.",
        "• Direct drone multispectral aerial data feeds."
    ]),
    ("Phase 3: Edge Deployment", [
        "• Quantized Granite-4-H-Small on local gateway devices.",
        "• Offline inference for zero-connectivity regions.",
        "• Autonomous smart-valve drip irrigation control.",
        "• Cooperative market price intelligence integration."
    ])
]
for i, (title, items) in enumerate(s13_data):
    x_val = float(0.8 + i * 4.0)
    create_card(slide_13, Inches(x_val), Inches(1.8), Inches(3.7), Inches(5.0), title, items)

# ==============================================================================
# SLIDES 14 & 15: Certificates (Indices 13 & 14)
# ==============================================================================
def embed_cert(slide_idx: int, title: str, base_name: str) -> None:
    slide = prs.slides[slide_idx]
    reset_slide_content(slide)
    set_slide_title(slide, title)
    img_file = find_image(base_name)
    if img_file:
        slide.shapes.add_picture(img_file, Inches(1.5), Inches(1.8), width=Inches(10.3))
    else:
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(1.8), Inches(10.3), Inches(4.8))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_CARD
        box.line.color.rgb = COLOR_BORDER
        tf_c = get_text_frame(box)
        if tf_c is not None:
            tf_c.text = f"[{base_name} image not found]"

embed_cert(13, "IBM SkillsBuild Completion Certificate", "skillsbuild_cert")
embed_cert(14, "IBM BOB Completion Certificate", "bob_cert")

prs.save(PPTX_PATH)
print("All 15 slides successfully synchronized with IBM template standard headings.")